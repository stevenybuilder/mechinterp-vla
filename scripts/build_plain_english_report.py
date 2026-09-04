#!/usr/bin/env python3
"""Build the five-page plain-English research report.

Design basis: the documents skill's narrative_proposal preset with an
editorial-cover header. The body size and margins are compacted slightly to
honor the user's explicit five-page limit while keeping the report readable.
"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "VLA Model Biology - Plain English Research Report.docx"
MATCHED = ROOT / "artifacts/pi05_matched_band_transform_2026-09-04_v1/summary.json"

NAVY = "102447"
BLUE = "1D4ED8"
TEAL = "0F766E"
ORANGE = "C2410C"
RED = "B42318"
INK = "172033"
MUTED = "536273"
PALE = "F4F6F9"
PALE_BLUE = "EAF1FF"
PALE_TEAL = "E9F7F4"
WHITE = "FFFFFF"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=55, start=90, bottom=55, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text: str, *, bold=False, color=INK, size=8.1) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)


def add_table(doc, headers, rows, widths=None, font_size=7.8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, color=WHITE, size=font_size)
        shade(table.rows[0].cells[i], NAVY)
        if widths:
            table.rows[0].cells[i].width = Inches(widths[i])
    set_repeat_table_header(table.rows[0])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value), size=font_size)
            if ridx % 2 == 1:
                shade(cells[i], PALE)
            if widths:
                cells[i].width = Inches(widths[i])
    return table


def add_field(run, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def add_para(doc, text: str = "", *, bold_lead: str | None = None, size=9.1, color=INK,
             align=None, after=3, before=0, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.06
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(NAVY)
        rest = p.add_run(text[len(bold_lead):])
        for run in (r, rest):
            run.font.name = "Calibri"
            run.font.size = Pt(size)
    else:
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor.from_string(color)
        r.italic = italic
    return p


def add_bullets(doc, items, *, size=8.8, compact=True):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.12)
        p.paragraph_format.space_after = Pt(1 if compact else 3)
        p.paragraph_format.line_spacing = 1.02
        r = p.add_run(item)
        r.font.name = "Calibri"
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor.from_string(INK)


def add_box(doc, title: str, body: str, *, fill=PALE_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(0.08)
    table.columns[1].width = Inches(6.95)
    shade(table.cell(0, 0), accent)
    shade(table.cell(0, 1), fill)
    set_cell_margins(table.cell(0, 0), 0, 0, 0, 0)
    cell = table.cell(0, 1)
    set_cell_margins(cell, 80, 120, 80, 120)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(10.2)
    r.font.color.rgb = RGBColor.from_string(accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.03
    r2 = p2.add_run(body)
    r2.font.name = "Calibri"
    r2.font.size = Pt(8.75)
    r2.font.color.rgb = RGBColor.from_string(INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_caption(doc, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.italic = True
    r.font.name = "Calibri"
    r.font.size = Pt(7.4)
    r.font.color.rgb = RGBColor.from_string(MUTED)


def add_page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)


def add_page_title(doc, kicker: str, title: str, deck: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(kicker.upper())
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(TEAL)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(3)
    r2 = p2.add_run(title)
    r2.bold = True
    r2.font.name = "Calibri"
    r2.font.size = Pt(17)
    r2.font.color.rgb = RGBColor.from_string(NAVY)
    if deck:
        add_para(doc, deck, size=9.1, color=MUTED, italic=True, after=5)


def matched_result_text() -> tuple[str, str]:
    if not MATCHED.exists():
        return (
            "Matched layer-6→8 test",
            "The final matched-versus-mismatched attention-plus-MLP panel was still running when this report was generated. Regenerate the report after its summary is archived.",
        )
    data = json.loads(MATCHED.read_text())
    cond = data["conditions"]
    con = data["contrasts"]
    mf = cond["matched_full"]["median_cell_progress_to_B"]
    mm = cond["mismatched_full_norm_matched"]["median_cell_progress_to_B"]
    ao = cond["matched_attention_only"]["median_cell_progress_to_B"]
    primary = con["matched_minus_mismatched_full"]
    reuse = con["mismatched_full_minus_random"]
    mlp = con["matched_full_minus_matched_attention_only"]
    flags = data["interpretation_flags"]
    return (
        "Matched layer-6→8 test",
        f"Across 12 directed prompt-pair cells, median A→B action-axis progress was {mf:.3f} for the matching full attention+MLP message, {mm:.3f} for the norm-matched message from another scene, and {ao:.3f} for the matching attention-only message. Other-scene full beat random in 12/12 cells (exact sign p=0.00049), and full beat attention-only in 12/12 (p=0.00049). Matching-scene messages were somewhat stronger than other-scene messages, but only 9/12 signs agreed (p=0.146): shared computation with suggestive scene modulation, not clean scene specificity.",
    )


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    section.header_distance = Inches(0.22)
    section.footer_distance = Inches(0.25)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9.1)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    for name, size, color in (("Title", 25, NAVY), ("Heading 1", 16, NAVY), ("Heading 2", 11.5, BLUE)):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("VLA MODEL BIOLOGY  /  EVIDENCE-CHECKED REPORT")
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(7.4)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Steven Yang  •  4 September 2026  •  ")
    r.font.name = "Calibri"
    r.font.size = Pt(7.3)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    add_field(r, "PAGE")
    r2 = p.add_run(" / 5")
    r2.font.name = "Calibri"
    r2.font.size = Pt(7.3)
    r2.font.color.rgb = RGBColor.from_string(MUTED)

    # Page 1 — application questions.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("CAUSAL INTERPRETABILITY OF A ROBOT POLICY")
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(TEAL)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("Where does π0.5 keep the instruction it is following?")
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    add_para(doc, "A five-page plain-English report on the project's lineage, experiments, evidence, failures, and safety relevance.", size=10.2, color=MUTED, italic=True, after=7)
    add_box(doc, "What question did you try to answer?", "When π0.5 turns a sentence and camera view into a movement, where does the instruction's causal influence go? We looked for a small internal object—a direction, token-local code, low-rank subspace, or compact pathway—that could be changed without replacing the rest of the scene and motor state.")
    add_box(doc, "Why is the question interesting?", "Language-model interpretability often starts with semantic tokens. A robot must mix language with pixels, pose, and control constraints before it acts. Its decision may live in a private coordinate system that is not readable as words. The setting also supplies unusually hard validation: same scene, different valid command, followed by action distance, first object touched, and task success.", fill=PALE_TEAL, accent=TEAL)
    add_box(doc, "What conclusions did you reach?", "The instruction is used during multimodal prefill. Later, its text-token copy remains readable but has little control over the tested Goal actions. Broad late image-position state can redirect behavior, including 18/20 closed-loop successes versus 0/20 for an early-layer control. Compact descriptions failed, but a full layer-6→8 update transferred across initial scenes and required the MLP sublayers. Decoding, causal use, and token provenance separate after modalities mix.", fill="FFF3E8", accent=ORANGE)
    add_table(doc, ["Headline", "Verified result", "Plain-English meaning"], [
        ("Causal handoff", "Text K/V R=0.0106; late image K/V R=0.8321", "The readable instruction and the action-controlling state are not in the same place."),
        ("Behavior", "Late live repair 18/20; early control 0/20", "The broad state changes complete robot behavior, not only one vector."),
        ("Compression", "Object-local 0/8; Sonar endpoints 0/12", "A selective, portable instruction feature was not isolated."),
    ], widths=[1.35, 2.15, 3.55], font_size=7.7)
    add_para(doc, "Main organism: π0.5 LIBERO checkpoint. Boundary test: OpenVLA-OFT. Evidence is scoped to the archived tasks and checkpoints; it is not a universal claim about all VLAs.", size=7.9, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, before=3)

    # Page 2 — lineage and conceptual diagrams.
    add_page_break(doc)
    add_page_title(doc, "Origins", "From Othello-style coordinates to physical action", "The project changed organisms, not its core concern: how to audit a decision that exists inside a learned latent state before it becomes visible behavior.")
    add_table(doc, ["Commit", "Project move", "Reason"], [
        ("1e080c1", "JEPA/world-model causal forensics", "Try to implant and locate a false physical rule."),
        ("ca3b272", "Pivot to VLA instruction locus", "The world-model organism lacked a validated behavioral effect; π0.5 + LIBERO had one."),
        ("32285db", "Reader-facing framing", "Make the same-scene instruction conflict and causal test legible."),
        ("2bfafb9", "Suite qualification", "Correct an overclaim after Goal and Object results differed."),
    ], widths=[0.8, 2.25, 4.0], font_size=7.4)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    paths = [
        Path("/Users/stevenyang/Downloads/othello.png"),
        Path("/Users/stevenyang/Downloads/mech interp wm.png"),
        Path("/Users/stevenyang/Downloads/representational geometry + superposition.png"),
    ]
    for cell, path in zip(table.rows[0].cells, paths):
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(2.28))
        set_cell_margins(cell, 30, 30, 30, 30)
    add_caption(doc, "Conceptual lineage (left to right): model-native coordinates; the longer LLM→VLA→world-model causal chain; directions, subspaces, superposition, and manifolds. These diagrams motivate tests; they are not result figures.")
    add_para(doc, "The Othello lesson was that a human label can be the wrong coordinate. “Black versus white” is less natural to the model than “mine versus theirs,” which flips with the turn. A VLA has more possible coordinate systems: camera-relative, gripper-relative, object-relative, goal-relative, and combinations of them. This is why a direction learned in one scene may fail in another without the information disappearing.", bold_lead="The Othello lesson", size=8.7)
    add_para(doc, "The first plan used a JEPA-style MetaWorld predictor with a deliberately false door rule. It never reached the intended interpretability experiment: the checkpoint and paired counterfactuals were not validated, and the planner produced 0/120 elite inaction samples. The pivot to a released VLA kept the latent-state safety question while giving us a working organism and objective behavior.", bold_lead="The first plan", size=8.7)
    add_para(doc, "The Plato's-cave intuition from Beyond Language Modeling also survives in a modest form. Text describes the physical world; it is not the world. In π0.5, a linguistic instruction becomes part of a multimodal state whose function is better described by the action it causes than by the token type at its storage position.", bold_lead="The Plato's-cave intuition", size=8.7)

    # Page 3 — experimental ladder and method.
    add_page_break(doc)
    add_page_title(doc, "What we did", "A ladder from behavior to causal mechanism", "Each step asked a stronger question. Positive internal geometry was never treated as enough without an action or rollout test.")
    doc.add_picture(str(ROOT / "figures/01_stimulus_intervention_behavior.png"), width=Inches(7.05))
    add_caption(doc, "Figure 1. Same pixels and robot state, two valid instructions. Broad late image-state replacement redirected closed-loop behavior; the early-layer control did not.")
    add_table(doc, ["Stage", "Direct test", "Main finding"], [
        ("Behavior", "Conflict the benchmark goal with another valid command", "Obedience became measurable by first touch and success."),
        ("Readout", "Probe each layer and inspect attention", "Instruction identity remained readable at text positions."),
        ("Causality", "Patch text, image, and other prefix state", "Late image state, not late text state, carried most Goal action control."),
        ("Timing", "Swap instruction before prefill; repair state after", "The instruction is used early, then its consequence is distributed."),
        ("Behavioral validation", "Repeat the patch at every replan", "Late repair completed 18/20 tasks; early control completed 0/20."),
        ("Compression", "Patch object tokens, position doses, directions, subspaces", "No small portable representation matched the whole-state effect."),
        ("Communication", "Block and rescue instruction→image and image→action edges", "The broad route was necessary; compact writer rescue failed."),
        ("L6–8 computation", "Curvature, then full/attention-only matched-scene tests", "The path bent; only the natural full-block update transferred across scenes."),
    ], widths=[1.0, 2.65, 3.4], font_size=7.05)
    add_para(doc, "Stimuli and controls.", bold_lead="Stimuli and controls.", size=8.4, before=3, after=1)
    add_bullets(doc, [
        "LIBERO Goal and Object tabletop scenes; A/B prompts were token-position matched where an intervention required it.",
        "The observation, robot state, action noise, and prompt formatting were held fixed; only the task instruction changed.",
        "Self-patches, wrong prompts, early/dead layers, equal-count random positions, norm-matched random directions, reversed signs, and distances to both clean endpoints guarded against silent no-ops and model damage.",
    ], size=8.15)

    # Page 4 — numbers and metrics.
    add_page_break(doc)
    add_page_title(doc, "Evidence", "What the main numbers actually measure", "The central result is a mismatch between where information can be read and where an intervention changes behavior.")
    doc.add_picture(str(ROOT / "figures/02_layers_attention_causality.png"), width=Inches(7.05))
    add_caption(doc, "Figure 2. Layerwise decoding, attention, and causal repair answer different questions. “Probe accuracy” is 10-way grouped cross-validated classification accuracy—not a correlation coefficient.")
    add_table(doc, ["Metric", "Definition", "Headline result"], [
        ("Probe accuracy", "10-way multinomial logistic regression; 5 folds grouped by initial state", "Text 1.000 at L0–16; image 0.113 at L0 → 0.993 at L1."),
        ("Repair R", "Progress from source toward donor action, with endpoint distance checks", "Text K/V 0.0106; late image K/V 0.8321."),
        ("Endpoint distance", "D_A and D_B are L2 distance divided by clean A–B distance", "Broad carrier D_B=0.2349; dose-256 D_B=0.9260; dose-512 D_B=0.2564."),
        ("Behavior", "First intended object touched; simulator task success", "Conflict 0/20; late repair 18/20; early repair 0/20."),
        ("Attention mass", "Action-query mass per prefix segment; both total and per-token", "Text/image 11.41× per token; image/text 4.99× in total."),
        ("Monitor F1", "Outcome classifier versus constant positive baseline", "Dev 0.909; held instruction 0.681; held scene 0.745; baseline 0.801."),
        ("Curvature/chord", "Midpoint deviation divided by endpoint-output chord", "Median 0.188; 12/12 cell medians ≥0.1."),
        ("Curvature action", "Action change after linearization divided by endpoint action distance", "Median 0.0973; 5/12 ≥0.1; preregistered fail."),
        ("Cross-model Dsrc", "Distance from patched action to donor, normalized by endpoints", "OpenVLA-OFT instruction 0.0101; image 0.9574."),
    ], widths=[1.12, 3.05, 2.88], font_size=6.85)
    add_para(doc, "Sample size discipline: the inferential unit was the directed prompt-pair or scene-condition cell. Repeated rollouts and five initial states were replicates inside a cell. Treating all raw rows as independent would manufacture confidence because many repeated rollouts are identical or highly dependent.", size=8.0, color=MUTED, before=3)

    # Page 5 — negative evidence, limitations, safety relevance, and final test.
    add_page_break(doc)
    add_page_title(doc, "What survived", "A real handoff—and strong evidence against the simple story", "The project is most informative when the positive and negative evidence are kept together.")
    title, body = matched_result_text()
    add_box(doc, title, body, fill="FFF3E8", accent=ORANGE)
    doc.add_picture(str(ROOT / "figures/07_matched_band_transform.png"), width=Inches(6.9))
    add_caption(doc, "Figure 3. A full instruction-induced update transfers across initial scenes within a prompt pair, while MLP-clamped attention-only updates are weak. Dots are 12 directed cell medians.")
    add_table(doc, ["Hypothesis", "Evidence against it"], [
        ("The instruction stays causally at text positions", "Text remains decodable, but post-prefill text repair is near zero; prefill swaps show it was used earlier."),
        ("The causal state is object-local", "Object-local replacement worked in 0/8 directions; 4–256 image positions were weak relative to all 512."),
        ("One portable direction is enough", "A static vector produced 5/5 correct first touches in one task but 0/5 successes and failed the two-task gate."),
        ("A low-rank subspace is the mechanism", "Rank-16 Sonar fit beat controls in 12/12 cells but reached 0/12 joint endpoints; donor-free repair was 0/30."),
        ("One compact writer seeds the state", "Full instruction→image blocking was A-like in 12/12, but block-0 rescue and fixed L6–8 rescue failed."),
        ("Observed nonlinearity controls action", "Curvature was real, yet removing it missed the frozen causal gate."),
    ], widths=[2.2, 4.85], font_size=7.25)
    add_para(doc, "Why this matters for safety.", bold_lead="Why this matters for safety.", size=8.7, before=4, after=1)
    add_para(doc, "A monitor can read a stale copy of the instruction while missing the state that controls motion. A broad edit can make a robot touch the intended object while damaging the later trajectory. As learned policies and world models move into robots and vehicles, interpretability needs external behavioral validation and collateral-damage tests—not only probes, attention maps, or attractive latent-space plots.", size=8.6)
    add_para(doc, "Limitations and what could have helped.", bold_lead="Limitations and what could have helped.", size=8.7, before=2, after=1)
    add_bullets(doc, [
        "The strongest repair is donor-assisted and replaces a large state. A learned conditional intervention needs true prompt-pair holdout and closed-loop preservation controls.",
        "The 8-edge reader result is development-only; it should not be called a confirmed circuit. The donor-free run is incomplete at 250/400 planned rows.",
        "Two checkpoints do not establish a VLA taxonomy. More models, training regimes, embodiments, and naturally occurring failure modes are needed.",
        "A future world-model study should begin with a released checkpoint and validated behavioral error before searching latent coordinates.",
    ], size=8.0)
    add_para(doc, "Bottom line: π0.5 contains a causally important language→multimodal-state→action handoff. We localized the broad state and changed behavior, but did not isolate a small reusable mechanism. That negative boundary is scientifically useful because it shows exactly where standard linear-feature and token-local stories stopped predicting causal control.", size=8.8, color=NAVY, before=3)

    core = doc.core_properties
    core.title = "Where does π0.5 keep the instruction it is following?"
    core.subject = "Plain-English evidence-checked report on VLA model biology"
    core.author = "Steven Yang"
    core.keywords = "mechanistic interpretability, VLA, pi0.5, LIBERO, model biology"
    core.comments = "Generated from the evidence-checked MATS archive; five-page narrative_proposal layout."
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
